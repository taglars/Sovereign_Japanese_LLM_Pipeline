"""The human rating form as an Excel workbook, for mail systems that block HTML."""

from __future__ import annotations

import html as H
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Tuple

ROOT = Path(__file__).resolve().parents[1]
FORM = ROOT / "data" / "rating_form.html"
OUT = ROOT / "data" / "rating_form.xlsx"

TEXT_WIDTH = 110          # column C, in Excel character units
JP_FONT = "Yu Gothic"
YELLOW = "FFF4C2"
GREY = "6B6B6B"


@dataclass
class Question:
    key: str
    title: str
    ask: str
    anchors: List[str]


@dataclass
class FormScript:
    id: str
    sections: List[Tuple[str, List[str]]]      # (label, visible lines)


def parse_form(path: Path = FORM) -> Tuple[List[FormScript], List[Question], str]:
    """Scripts, questions and the instruction text, as the HTML form has them."""
    src = path.read_text(encoding="utf-8")
    js = "\n".join(re.findall(r"(?s)<script[^>]*>(.*?)</script>", src))
    dec = json.JSONDecoder()
    start = js.find("const DATA = ") + len("const DATA = ")
    data, end = dec.raw_decode(js[start:])
    rest = js[start + end:]
    qhtml, _ = dec.raw_decode(rest[rest.find("QUESTIONS = ") + len("QUESTIONS = "):])

    questions = []
    for fs in re.findall(r'(?s)<fieldset class="q">(.*?)</fieldset>', qhtml):
        strip = lambda s: H.unescape(re.sub(r"<[^>]+>", "", s)).strip()  # noqa: E731
        questions.append(Question(
            key=re.search(r'name="(\w+)"', fs).group(1),
            title=strip(re.search(r"<legend>(.*?)</legend>", fs).group(1)),
            ask=re.sub(r"\s+\n|\n\s+", "\n", strip(
                re.search(r'(?s)<p class="ask">(.*?)</p>', fs).group(1))),
            anchors=[strip(a) for a in re.findall(r"(?s)<li>(.*?)</li>", fs)]))

    scripts = []
    for item in data:
        sections = []
        for label, body in re.findall(
                r'(?s)<section><h4>(.*?)</h4><div class="jp">(.*?)</div></section>',
                item["html"]):
            # What the rater saw: the escaped markup rendered as text, <br> as a line
            # break. Unescaping once reproduces it exactly.
            lines = [H.unescape(line) for line in body.split("<br>")]
            sections.append((H.unescape(label), [l for l in lines if l.strip()]))
        scripts.append(FormScript(id=item["id"], sections=sections))

    intro = re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", " ", re.search(
        r"(?s)Read each script through.*?so do not try to work out which\.", src).group(0))))
    return scripts, questions, intro.strip()


def _height(text: str, width: int = TEXT_WIDTH, per_line: float = 16.0) -> float:
    """Row height for wrapped text. Excel does not reliably auto-fit rows written by
    openpyxl, and a squashed row hides most of a script line.
    """
    lines = 0
    for para in text.split("\n"):
        units = sum(2 if ord(ch) > 0x2E80 else 1 for ch in para)
        lines += max(1, math.ceil(units / (width * 0.92)))
    return min(409.0, max(16.0, lines * per_line + 4))


def build(out: Path = OUT, form: Path = FORM) -> Path:
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    scripts, questions, intro = parse_form(form)
    wb = openpyxl.Workbook()
    start = wb.active
    start.title = "Start"

    wrap = Alignment(wrap_text=True, vertical="top")
    answer_fill = PatternFill("solid", fgColor=YELLOW)
    thin = Side(style="thin", color="B5A642")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    unlocked = Protection(locked=False)

    def setup(ws):
        ws.column_dimensions["A"].hidden = True
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = TEXT_WIDTH
        ws.sheet_view.showGridLines = False
        # Protected without a password: script text cannot be overwritten by accident, the
        # yellow cells stay editable, rows and columns resizable.
        ws.protection.sheet = True
        ws.protection.formatRows = False
        ws.protection.formatColumns = False

    # -- per-script sheets ---------------------------------------------------
    answer_rows = {}
    for n, script in enumerate(scripts, 1):
        ws = wb.create_sheet(f"Script {n}")
        setup(ws)
        ws["A1"] = f"script_id:{script.id}"
        ws["B1"] = f"Script {n} of {len(scripts)}"
        ws["B1"].font = Font(bold=True, size=14)
        row = 3
        for label, lines in script.sections:
            ws.cell(row=row, column=2, value=label).font = Font(size=9, color=GREY, bold=True)
            row += 1
            for line in lines:
                cell = ws.cell(row=row, column=3, value=line)
                cell.font = Font(name=JP_FONT, size=11)
                cell.alignment = wrap
                ws.row_dimensions[row].height = _height(line)
                row += 1
            row += 1

        row += 1
        ws.cell(row=row, column=2, value="Your ratings").font = Font(bold=True, size=13)
        ws.cell(row=row, column=3, value="Enter one whole number from 1 to 10 in each "
                "yellow box. The descriptions give examples for some points on the "
                "scale; the numbers in between are just as valid.").font = Font(
                    italic=True, color=GREY)
        ws.cell(row=row, column=3).alignment = wrap
        row += 2

        first = row
        dv = DataValidation(type="list", formula1='"1,2,3,4,5,6,7,8,9,10"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="1 to 10 only",
                            error="Please enter one whole number from 1 to 10.")
        ws.add_data_validation(dv)
        for q in questions:
            ws.cell(row=row, column=1, value=q.key)
            ans = ws.cell(row=row, column=2)
            ans.fill, ans.border, ans.protection = answer_fill, box, unlocked
            ans.font = Font(bold=True, size=14)
            ans.alignment = Alignment(horizontal="center", vertical="center")
            dv.add(ans)
            body = q.ask + "\n" + "\n".join(q.anchors)
            ws.cell(row=row, column=3).value = CellRichText(
                TextBlock(InlineFont(b=True, sz=12), q.title + "\n"),
                TextBlock(InlineFont(sz=11), body))
            ws.cell(row=row, column=3).alignment = wrap
            ws.row_dimensions[row].height = _height(q.title + "\n" + body)
            row += 2
        answer_rows[n] = (first, row - 2)

        ws.cell(row=row, column=1, value="comments")
        ws.cell(row=row, column=2, value="Anything worth noting (optional)").alignment = wrap
        note = ws.cell(row=row, column=3)
        note.fill, note.border, note.protection, note.alignment = answer_fill, box, unlocked, wrap
        ws.row_dimensions[row].height = 80

    # -- start sheet -----------------------------------------------------------
    setup(start)
    start["B1"] = f"Sales script rating, {len(scripts)} scripts"
    start["B1"].font = Font(bold=True, size=16)
    start["C3"] = intro
    start["C3"].alignment = wrap
    start.row_dimensions[3].height = _height(intro)
    start["C4"] = ("Each script is on its own tab at the bottom: Script 1 to Script "
                   f"{len(scripts)}. Under each script are six questions; enter your "
                   "answer in the yellow box beside each one.")
    start["C4"].alignment = wrap
    start.row_dimensions[4].height = 36

    start["A6"] = "rater"
    start["B6"] = "Your name or initials:"
    start["B6"].alignment = wrap
    name = start["C6"]
    name.fill, name.border, name.protection = answer_fill, box, unlocked
    name.font = Font(size=12)
    start.row_dimensions[6].height = 22

    start["B8"] = "When you are done"
    start["B8"].font = Font(bold=True)
    start["C8"] = "Save this file and send it back by email. Nothing else is needed."
    start["C8"].alignment = wrap

    start["B10"] = "Progress"
    start["B10"].font = Font(bold=True)
    start["C10"] = "Fills in as you answer. Every script needs all six questions."
    start["C10"].font = Font(italic=True, color=GREY)
    r = 11
    totals = []
    for n, (a, b) in answer_rows.items():
        rng = f"'Script {n}'!B{a}:B{b}"
        start.cell(row=r, column=2, value=f"Script {n}")
        start.cell(row=r, column=3, value=f'=COUNT({rng})&" of {len(questions)} answered"')
        totals.append(f"COUNT({rng})")
        r += 1
    start.cell(row=r + 1, column=2, value="All scripts").font = Font(bold=True)
    start.cell(row=r + 1, column=3,
               value=f'=({"+".join(totals)})&" of {len(questions) * len(scripts)} answered"'
               ).font = Font(bold=True)

    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def read(path: Path) -> List[dict]:
    """A returned workbook -> rows shaped like the HTML form's CSV export: rater,
    script_id, the six *_score keys, comments.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rater = ""
    for row in wb["Start"].iter_rows(min_col=1, max_col=3):
        if row[0].value == "rater":
            rater = str(row[2].value or "").strip()
    rows = []
    for ws in wb.worksheets:
        sid, got = "", {}
        for cells in ws.iter_rows(min_col=1, max_col=3):
            key = str(cells[0].value or "")
            if key.startswith("script_id:"):
                sid = key.split(":", 1)[1]
            elif key == "comments":
                got["comments"] = cells[2].value
            elif key.endswith("_score"):
                got[key] = cells[1].value
        if sid:
            rows.append({"rater": rater, "script_id": sid, **got})
    return rows


def is_rating_workbook(path: Path) -> bool:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        return "Start" in wb.sheetnames and any(n.startswith("Script ") for n in wb.sheetnames)
    except Exception:                                          # noqa: BLE001
        return False
