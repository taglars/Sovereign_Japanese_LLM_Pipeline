"""The Japanese prompt table, out to a workbook for review and back again."""

from __future__ import annotations

import collections
import json
import re
from pathlib import Path
from typing import Dict, List

from . import templates as T
from pipeline.common.prompts.registry import Fidelity, get as get_prompts  # noqa: E402
from pipeline.common.tasks import RecommendationType  # noqa: E402

SHEET = T.TABLE_PATH.with_name("review_batch_prod_ja.xlsx")
_TAG = re.compile(r"</?[a-z]+>")


def usage() -> Dict[str, List[str]]:
    """Which prompts each line appears in, so the reviewer sees its context."""
    prompts = get_prompts(Fidelity.BATCH_PROD)
    where: Dict[str, List[str]] = collections.defaultdict(list)
    for family in RecommendationType:
        for spec in T.stage_graph(prompts, family):
            label = f"{family.value}/{spec.task}" + (f"/{spec.section}" if spec.section else "")
            for part in T._SENTINEL.split(T.template(prompts, spec))[0::2]:
                for _, content, _ in T._units(part):
                    if content and label not in where[content]:
                        where[content].append(label)
    for text in T.lease_units():
        for _, content, _ in T._units(text):
            if content:
                where[content].append("lease text (stored as data)")
    where[T.LANG_EN].append("lang directive (every prompt)")
    return where


def write_sheet(path: Path = SHEET) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    rows = json.loads(T.TABLE_PATH.read_text(encoding="utf-8"))
    where = usage()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "translations"
    header = ["id", "used in", "english (production)", "japanese (draft)",
              "correction (leave empty if fine)", "note"]
    ws.append(header)
    for row in rows:
        places = where.get(row["en"], [])
        shown = ", ".join(places[:4]) + (f" (+{len(places) - 4})" if len(places) > 4 else "")
        ws.append([row["id"], shown, row["en"], row["ja"], "", ""])
    widths = [7, 28, 70, 70, 50, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEEEEE")
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    guide = wb.create_sheet("how to review")
    for line in [
        "Review the Japanese draft of each production prompt line.",
        "Write a correction ONLY where the draft is wrong or unnatural. Empty = keep.",
        "Keep exactly: numbering (1. a. i.), **bold markers**, HTML tags like <li>, "
        "and the placeholder {term}.",
        # product names withheld from this copy (see README)
        "Product and feature names in katakana or English ([names withheld]) "
        "stay as in the draft unless wrong.",
        "Quoted examples of bad output are translated on purpose: the model writes "
        "Japanese, so its examples are Japanese.",
        "Known production quirks are kept, not fixed, e.g. the TCO prompt saying "
        "'range', or 'mileage lease' mapped to パワーリース. They are what the "
        "company sent.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def apply(path: Path) -> Dict[str, int]:
    import openpyxl
    rows = json.loads(T.TABLE_PATH.read_text(encoding="utf-8"))
    by_id = {r["id"]: r for r in rows}
    ws = openpyxl.load_workbook(path, data_only=True)["translations"]
    header = [str(c.value or "") for c in ws[1]]
    col = {name: header.index(name) for name in header}
    changed, problems = 0, []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        rid = raw[col["id"]]
        fix = str(raw[col["correction (leave empty if fine)"]] or "").strip()
        if not rid or not fix or rid not in by_id:
            continue
        en = by_id[rid]["en"]
        if ("{term}" in en) != ("{term}" in fix):
            problems.append(f"{rid}: {{term}} placeholder lost or added")
        if sorted(_TAG.findall(en)) != sorted(_TAG.findall(fix)):
            problems.append(f"{rid}: HTML tags differ from the English line")
        if en.count("**") != fix.count("**"):
            problems.append(f"{rid}: ** bold markers differ from the English line")
        if fix != by_id[rid]["ja"]:
            by_id[rid]["ja"] = fix
            by_id[rid]["reviewed"] = True
            changed += 1
    if problems:
        raise ValueError("review not applied, fix these first:\n  " + "\n  ".join(problems))
    T.TABLE_PATH.write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    T.load_table()                                    # the same checks the job runs
    return {"changed": changed, "lines": len(rows)}
