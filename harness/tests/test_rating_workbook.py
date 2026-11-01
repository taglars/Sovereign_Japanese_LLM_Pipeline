"""The Excel rating form: same stimulus as the HTML form, and it reads back."""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from jv import human                                               # noqa: E402
from jv import rating_workbook as RW                               # noqa: E402
from jv.rubric import SCORED_KEYS                                  # noqa: E402


def test_same_scripts_and_questions_as_the_html_form():
    scripts, questions, intro = RW.parse_form()
    assert [s.id for s in scripts] == [f"S{i:02d}" for i in range(1, 13)]
    assert [q.key for q in questions] == list(SCORED_KEYS)
    assert all(len(q.anchors) == 4 for q in questions)
    assert intro.startswith("Read each script through")


def test_filled_workbook_reads_back_exactly():
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        path = RW.build(Path(tmp) / "form.xlsx")
        wb = openpyxl.load_workbook(path)
        wb["Start"]["C6"] = "  A. Rater "
        want = {}
        for i, name in enumerate(wb.sheetnames[1:]):
            sid = None
            for row in wb[name].iter_rows(min_col=1, max_col=3):
                key = str(row[0].value or "")
                if key.startswith("script_id:"):
                    sid = key.split(":", 1)[1]
                elif key in SCORED_KEYS:
                    # LibreOffice may store a dropdown choice as text
                    value = (i + len(key)) % 10 + 1
                    row[1].value = str(value) if i % 2 else value
                    want[(sid, key)] = value
        wb.save(path)
        records = human.read_file(path, {})
        assert {r.judge for r in records} == {"human:A. Rater"}
        assert {(r.script_id, k): v for r in records for k, v in r.scores.items()} == want


def test_a_partly_filled_workbook_keeps_what_is_there():
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        path = RW.build(Path(tmp) / "form.xlsx")
        wb = openpyxl.load_workbook(path)
        for row in wb["Script 3"].iter_rows(min_col=1, max_col=3):
            if row[0].value == "keigo_score":
                row[1].value = 4
        wb.save(path)
        records = human.read_file(path, {})
        assert len(records) == 1 and records[0].scores == {"keigo_score": 4}


if __name__ == "__main__":
    # The form holds the twelve blind scripts, which are withheld from this copy.
    if not RW.FORM.exists():
        print("rating workbook: skipped, data/rating_form.html is withheld (see README)")
        raise SystemExit(0)
    failures = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"  ok  {name}")
            except Exception as exc:                                # noqa: BLE001
                failures += 1
                print(f"  FAIL {name}: {exc!r}"[:300])
    print("rating workbook: " + ("all checks pass" if not failures else f"{failures} FAILED"))
    raise SystemExit(1 if failures else 0)
